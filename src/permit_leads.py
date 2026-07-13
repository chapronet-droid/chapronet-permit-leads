from __future__ import annotations

import argparse
import json
import logging
import os
import smtplib
import sys
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from email.message import EmailMessage
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus, urlencode

import pandas as pd
import requests
import yaml
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


LOGGER = logging.getLogger("chapronet_permit_leads")


COLUMN_ALIASES = {
    "permit_number": ["permit_", "permit_number", "permit_no"],
    "permit_status": ["permit_status", "status"],
    "permit_milestone": ["permit_milestone", "milestone"],
    "permit_type": ["permit_type"],
    "review_type": ["review_type"],
    "application_start_date": ["application_start_date"],
    "issue_date": ["issue_date"],
    "street_number": ["street_number"],
    "street_direction": ["street_direction"],
    "street_name": ["street_name"],
    "work_type": ["work_type"],
    "work_description": ["work_description"],
    "reported_cost": ["reported_cost", "estimated_cost"],
    "contact_1_type": ["contact_1_type"],
    "contact_1_name": ["contact_1_name"],
    "contact_1_city": ["contact_1_city"],
    "contact_1_state": ["contact_1_state"],
    "contact_1_zipcode": ["contact_1_zipcode", "contact_1_zip"],
    "contact_2_type": ["contact_2_type"],
    "contact_2_name": ["contact_2_name"],
    "contact_3_type": ["contact_3_type"],
    "contact_3_name": ["contact_3_name"],
    "ward": ["ward"],
    "community_area": ["community_area"],
    "zip_code": ["zip_code", "zipcode"],
    "latitude": ["latitude"],
    "longitude": ["longitude"],
    "location": ["location"],
}


@dataclass(frozen=True)
class Paths:
    project_root: Path
    output_dir: Path
    state_dir: Path
    seen_file: Path


def load_config(config_path: Path) -> dict[str, Any]:
    with config_path.open("r", encoding="utf-8") as handle:
        config = yaml.safe_load(handle)
    if not isinstance(config, dict):
        raise ValueError("config.yaml must contain a mapping.")
    return config


def setup_paths(project_root: Path, config: dict[str, Any]) -> Paths:
    output_dir = project_root / config["output"]["directory"]
    state_dir = project_root / config["state"]["directory"]
    output_dir.mkdir(parents=True, exist_ok=True)
    state_dir.mkdir(parents=True, exist_ok=True)
    return Paths(
        project_root=project_root,
        output_dir=output_dir,
        state_dir=state_dir,
        seen_file=state_dir / config["state"]["seen_permits_filename"],
    )


def api_url(config: dict[str, Any]) -> str:
    src = config["source"]
    return f"https://{src['domain']}/resource/{src['dataset_id']}.json"


def build_where(config: dict[str, Any], start_date: str, end_date: str) -> str:
    """Build a Socrata filter using only columns present in the permit dataset.

    The Chicago Building Permits dataset does not expose a project ZIP-code
    column. Geographic targeting is therefore applied locally after download
    using ward/community_area, rather than sending an invalid zip_code clause.
    """
    clauses = [
        f"issue_date >= '{start_date}T00:00:00.000'",
        f"issue_date < '{end_date}T00:00:00.000'",
    ]

    # Cost and geography filters are applied locally. Keeping the API query to
    # date fields only avoids failures caused by dataset schema/type changes.
    return " AND ".join(clauses)


def fetch_permits(
    config: dict[str, Any],
    start_date: str,
    end_date: str,
    session: requests.Session | None = None,
) -> list[dict[str, Any]]:
    src = config["source"]
    page_size = int(src.get("page_size", 1000))
    max_pages = int(src.get("max_pages", 50))
    timeout = int(src.get("timeout_seconds", 45))
    url = api_url(config)
    where = build_where(config, start_date, end_date)

    headers = {"User-Agent": "ChaproNet-Permit-Leads/1.0"}
    token = os.getenv("SOCRATA_APP_TOKEN", "").strip()
    if token:
        headers["X-App-Token"] = token

    client = session or requests.Session()
    records: list[dict[str, Any]] = []

    for page in range(max_pages):
        params = {
            "$where": where,
            "$order": "issue_date DESC, permit_ DESC",
            "$limit": page_size,
            "$offset": page * page_size,
        }
        for attempt in range(4):
            try:
                response = client.get(url, params=params, headers=headers, timeout=timeout)
                if not response.ok:
                    detail = response.text.strip().replace("\n", " ")[:1000]
                    raise requests.HTTPError(
                        f"{response.status_code} {response.reason}. API response: {detail}",
                        response=response,
                    )
                batch = response.json()
                if not isinstance(batch, list):
                    raise RuntimeError(f"Unexpected API response: {batch!r}")
                records.extend(batch)
                LOGGER.info("Fetched %s permits (page %s).", len(batch), page + 1)
                break
            except (requests.RequestException, ValueError) as exc:
                if attempt == 3:
                    raise RuntimeError(f"Chicago permit API request failed: {exc}") from exc
                sleep_seconds = 2 ** attempt
                LOGGER.warning("API attempt failed; retrying in %ss: %s", sleep_seconds, exc)
                time.sleep(sleep_seconds)

        if len(batch) < page_size:
            break
    else:
        LOGGER.warning("Reached max_pages=%s. Increase it if results were truncated.", max_pages)

    return records


def first_present(record: dict[str, Any], aliases: list[str]) -> Any:
    for alias in aliases:
        value = record.get(alias)
        if value not in (None, ""):
            return value
    return None


def normalize_record(record: dict[str, Any]) -> dict[str, Any]:
    result = {
        output_name: first_present(record, aliases)
        for output_name, aliases in COLUMN_ALIASES.items()
    }

    address_parts = [
        result.get("street_number"),
        result.get("street_direction"),
        result.get("street_name"),
    ]
    result["address"] = " ".join(str(x).strip() for x in address_parts if x not in (None, ""))
    result["reported_cost"] = pd.to_numeric(result.get("reported_cost"), errors="coerce")
    result["issue_date"] = pd.to_datetime(result.get("issue_date"), errors="coerce")
    result["application_start_date"] = pd.to_datetime(
        result.get("application_start_date"), errors="coerce"
    )

    permit_number = str(result.get("permit_number") or "").strip()
    result["source_url"] = (
        "https://data.cityofchicago.org/Buildings/Building-Permits/ydr8-5enu/"
        f"explore/query/SELECT%20*%20WHERE%20permit_%3D%27{permit_number}%27"
        if permit_number
        else "https://data.cityofchicago.org/Buildings/Building-Permits/ydr8-5enu"
    )

    contacts = []
    for i in range(1, 4):
        ctype = result.get(f"contact_{i}_type")
        cname = result.get(f"contact_{i}_name")
        if cname:
            contacts.append(f"{ctype or 'Contact'}: {cname}")
    result["permit_contacts"] = " | ".join(contacts)

    address = result.get("address", "")
    contact_names = " ".join(str(result.get(f"contact_{i}_name") or "") for i in range(1, 4)).strip()
    search_base = f"{address} Chicago IL {contact_names}".strip()
    result["google_maps_url"] = f"https://www.google.com/maps/search/?api=1&query={quote_plus(address + ' Chicago IL')}" if address else ""
    result["owner_research_url"] = f"https://www.google.com/search?q={quote_plus(address + ' Chicago property owner')}" if address else ""
    result["gc_research_url"] = f"https://www.google.com/search?q={quote_plus(search_base + ' general contractor phone email')}" if search_base else ""
    result["company_research_url"] = f"https://www.google.com/search?q={quote_plus(contact_names + ' Chicago company phone email')}" if contact_names else ""
    result["cook_county_search_url"] = "https://www.cookcountyassessor.com/address-search"
    result["outreach_status"] = "Not Researched"
    result["assigned_to"] = ""
    result["primary_contact"] = ""
    result["phone"] = ""
    result["email"] = ""
    result["company"] = ""
    result["next_follow_up"] = ""
    result["notes"] = ""
    return result


def contains_any(text: str, terms: list[str]) -> tuple[bool, list[str]]:
    text_lower = text.lower()
    hits = [term for term in terms if term.lower() in text_lower]
    return bool(hits), hits


def score_record(row: pd.Series, config: dict[str, Any], today: date) -> tuple[int, str]:
    scoring = config["scoring"]
    keywords = config["keywords"]
    score = 0
    reasons: list[str] = []

    cost = row.get("reported_cost")
    if pd.notna(cost):
        if cost >= scoring["high_value_cost"]:
            score += scoring["high_value_points"]
            reasons.append(f"project cost ${cost:,.0f}")
        elif cost >= scoring["medium_value_cost"]:
            score += scoring["medium_value_points"]
            reasons.append(f"project cost ${cost:,.0f}")
        else:
            score += scoring["base_cost_points"]

    permit_type = str(row.get("permit_type") or "")
    combined = " ".join(
        str(row.get(field) or "")
        for field in ("permit_type", "work_type", "work_description")
    )

    if "NEW CONSTRUCTION" in permit_type.upper():
        score += scoring["new_construction_points"]
        reasons.append("new construction")
    if "RENOVATION" in permit_type.upper() or "ALTERATION" in permit_type.upper():
        score += scoring["renovation_points"]
        reasons.append("renovation/alteration")

    for category, point_key, label in [
        ("commercial", "commercial_keyword_points", "commercial use"),
        ("multifamily", "multifamily_keyword_points", "multifamily"),
        ("low_voltage", "low_voltage_keyword_points", "direct low-voltage signal"),
        ("target_work", "target_work_keyword_points", "target work"),
    ]:
        matched, hits = contains_any(combined, keywords.get(category, []))
        if matched:
            score += scoring[point_key]
            reasons.append(f"{label}: {', '.join(hits[:3])}")

    issue_date = row.get("issue_date")
    if pd.notna(issue_date):
        age = (today - issue_date.date()).days
        if age <= 2:
            score += scoring["recent_2_days_points"]
            reasons.append("issued within 2 days")
        elif age <= 7:
            score += scoring["recent_7_days_points"]
            reasons.append("issued within 7 days")

    geo = config.get("geography", {})
    target_wards = {_normalized_geo_value(x) for x in geo.get("wards", [])}
    target_areas = {_normalized_geo_value(x) for x in geo.get("community_areas", [])}
    in_target_area = (
        _normalized_geo_value(row.get("ward")) in target_wards
        or _normalized_geo_value(row.get("community_area")) in target_areas
    )
    if in_target_area:
        score += scoring["south_side_points"]
        reasons.append("target South Side area")

    return int(score), "; ".join(reasons)


def _normalized_geo_value(value: Any) -> str:
    if value in (None, "") or pd.isna(value):
        return ""
    text = str(value).strip()
    try:
        return str(int(float(text)))
    except ValueError:
        return text


def apply_geography_filter(df: pd.DataFrame, config: dict[str, Any]) -> pd.DataFrame:
    """Apply South Side targeting with dataset-supported geography fields."""
    geo = config.get("geography", {})
    target_wards = {_normalized_geo_value(x) for x in geo.get("wards", [])}
    target_areas = {_normalized_geo_value(x) for x in geo.get("community_areas", [])}

    masks = []
    if target_wards and "ward" in df.columns:
        masks.append(df["ward"].map(_normalized_geo_value).isin(target_wards))
    if target_areas and "community_area" in df.columns:
        masks.append(df["community_area"].map(_normalized_geo_value).isin(target_areas))

    if not masks:
        return df

    combined = masks[0]
    for mask in masks[1:]:
        combined = combined | mask
    return df[combined].copy()


def prepare_dataframe(records: list[dict[str, Any]], config: dict[str, Any]) -> pd.DataFrame:
    if not records:
        return pd.DataFrame()

    normalized = [normalize_record(record) for record in records]
    df = pd.DataFrame(normalized)
    df = apply_geography_filter(df, config)

    minimum_cost = float(config.get("filters", {}).get("minimum_reported_cost", 0) or 0)
    if minimum_cost and "reported_cost" in df.columns:
        df = df[(df["reported_cost"].isna()) | (df["reported_cost"] >= minimum_cost)].copy()

    excluded_statuses = {
        str(x).upper() for x in config.get("filters", {}).get("exclude_permit_statuses", [])
    }
    excluded_types = {
        str(x).upper() for x in config.get("filters", {}).get("exclude_permit_types", [])
    }

    if "permit_status" in df.columns and excluded_statuses:
        df = df[
            ~df["permit_status"].fillna("").str.upper().isin(excluded_statuses)
        ].copy()
    if "permit_type" in df.columns and excluded_types:
        df = df[
            ~df["permit_type"].fillna("").str.upper().isin(excluded_types)
        ].copy()

    today = datetime.now(timezone.utc).date()
    scored = df.apply(lambda row: score_record(row, config, today), axis=1)
    df["lead_score"] = [item[0] for item in scored]
    df["score_reasons"] = [item[1] for item in scored]

    df["priority"] = pd.cut(
        df["lead_score"],
        bins=[-1, 34, 54, 74, float("inf")],
        labels=["Research", "Warm", "Hot", "Immediate"],
    ).astype(str)

    df["recommended_action"] = df["priority"].map(
        {
            "Immediate": "Research owner/GC today; call and email",
            "Hot": "Research owner/GC within 24 hours",
            "Warm": "Add to outreach sequence",
            "Research": "Review manually; monitor project",
        }
    )

    order = [
        "lead_score", "priority", "issue_date", "permit_number", "permit_status",
        "permit_type", "work_type", "address", "zip_code", "ward", "community_area",
        "reported_cost", "work_description", "permit_contacts", "contact_1_type",
        "contact_1_name", "contact_1_city", "contact_1_state", "contact_1_zipcode",
        "contact_2_type", "contact_2_name", "contact_3_type", "contact_3_name",
        "score_reasons", "recommended_action", "permit_contacts",
        "primary_contact", "company", "phone", "email", "outreach_status",
        "assigned_to", "next_follow_up", "notes", "source_url",
        "google_maps_url", "owner_research_url", "gc_research_url",
        "company_research_url", "cook_county_search_url", "latitude", "longitude",
        "review_type", "permit_milestone", "application_start_date",
    ]
    existing = [column for column in order if column in df.columns]
    extras = [column for column in df.columns if column not in existing]
    df = df[existing + extras]
    return df.sort_values(
        by=["lead_score", "reported_cost", "issue_date"],
        ascending=[False, False, False],
        na_position="last",
    ).reset_index(drop=True)


def load_seen(path: Path) -> set[str]:
    if not path.exists():
        return set()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return {str(item) for item in data}
    except (json.JSONDecodeError, OSError):
        LOGGER.warning("Could not read state file. Starting with empty state.")
        return set()


def save_seen(path: Path, permit_numbers: set[str]) -> None:
    trimmed = sorted(permit_numbers)[-100000:]
    path.write_text(json.dumps(trimmed, indent=2), encoding="utf-8")


def style_excel(path: Path) -> None:
    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in worksheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[1].height = 28

        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            values = [str(cell.value or "") for cell in column_cells[:200]]
            width = min(max(max(map(len, values), default=10) + 2, 12), 55)
            worksheet.column_dimensions[column_letter].width = width

        headers = {cell.value: cell.column for cell in worksheet[1]}
        if "work_description" in headers:
            col = headers["work_description"]
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
        if "score_reasons" in headers:
            col = headers["score_reasons"]
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
        for url_header in ["source_url", "google_maps_url", "owner_research_url", "gc_research_url", "company_research_url", "cook_county_search_url"]:
            if url_header in headers:
                col = headers[url_header]
                for row in range(2, worksheet.max_row + 1):
                    cell = worksheet.cell(row, col)
                    if cell.value:
                        cell.hyperlink = str(cell.value)
                        cell.style = "Hyperlink"
        if "lead_score" in headers and worksheet.max_row >= 2:
            score_letter = get_column_letter(headers["lead_score"])
            worksheet.conditional_formatting.add(
                f"{score_letter}2:{score_letter}{worksheet.max_row}",
                ColorScaleRule(
                    start_type="min", start_color="F8696B",
                    mid_type="percentile", mid_value=50, mid_color="FFEB84",
                    end_type="max", end_color="63BE7B",
                ),
            )
    workbook.save(path)


def export_reports(df: pd.DataFrame, new_df: pd.DataFrame, config: dict[str, Any], paths: Paths) -> dict[str, Path]:
    output_cfg = config["output"]
    excel_path = paths.output_dir / output_cfg["excel_filename"]
    csv_path = paths.output_dir / output_cfg["csv_filename"]
    top_path = paths.output_dir / output_cfg["top_leads_filename"]
    jobber_path = paths.output_dir / output_cfg.get("jobber_filename", "jobber_ready_leads.csv")

    df.to_csv(csv_path, index=False)
    threshold = int(output_cfg.get("minimum_score_for_top_lead", 35))
    maximum = int(output_cfg.get("maximum_top_leads", 100))
    top_df = df[df["lead_score"] >= threshold].head(maximum) if not df.empty else df
    top_df.to_csv(top_path, index=False)

    jobber = pd.DataFrame()
    if not new_df.empty:
        jobber["First name"] = "Permit"
        jobber["Last name"] = new_df["permit_number"].fillna("").astype(str)
        jobber["Company name"] = new_df.get("contact_1_name", "")
        jobber["Phone number"] = new_df.get("phone", "")
        jobber["Email address"] = new_df.get("email", "")
        jobber["Street 1"] = new_df.get("address", "")
        jobber["City"] = "Chicago"
        jobber["State"] = "IL"
        jobber["Zip code"] = new_df.get("zip_code", "")
        jobber["Notes"] = (
            "Permit " + new_df["permit_number"].fillna("").astype(str)
            + " | Score " + new_df["lead_score"].fillna(0).astype(str)
            + " | " + new_df["work_description"].fillna("").astype(str)
            + " | " + new_df["source_url"].fillna("").astype(str)
        )
    jobber.to_csv(jobber_path, index=False)

    summary = pd.DataFrame(
        [
            {"Metric": "Run time", "Value": datetime.now().astimezone().isoformat(timespec="seconds")},
            {"Metric": "Permits returned", "Value": len(df)},
            {"Metric": "New permits", "Value": len(new_df)},
            {"Metric": "Immediate leads", "Value": int((df.get("priority") == "Immediate").sum()) if not df.empty else 0},
            {"Metric": "Hot leads", "Value": int((df.get("priority") == "Hot").sum()) if not df.empty else 0},
            {"Metric": "Top-lead threshold", "Value": threshold},
        ]
    )

    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="All Leads", index=False)
        new_df.to_excel(writer, sheet_name="New Since Last Run", index=False)
        top_df.to_excel(writer, sheet_name="Top Leads", index=False)
        new_df.to_excel(writer, sheet_name="Outreach Queue", index=False)
        summary.to_excel(writer, sheet_name="Run Summary", index=False)
    style_excel(excel_path)

    return {"excel": excel_path, "csv": csv_path, "top": top_path, "jobber": jobber_path}


def email_report(files: dict[str, Path], new_df: pd.DataFrame) -> None:
    if os.getenv("EMAIL_ENABLED", "false").lower() != "true":
        return
    if new_df.empty and os.getenv("EMAIL_ONLY_WHEN_NEW", "true").lower() == "true":
        LOGGER.info("Email skipped because there are no new leads.")
        return

    required = ["SMTP_HOST", "SMTP_PORT", "SMTP_USERNAME", "SMTP_PASSWORD", "EMAIL_FROM", "EMAIL_TO"]
    missing = [name for name in required if not os.getenv(name)]
    if missing:
        raise RuntimeError(f"Email is enabled but these settings are missing: {', '.join(missing)}")

    message = EmailMessage()
    message["Subject"] = f"ChaproNet permit leads: {len(new_df)} new"
    message["From"] = os.environ["EMAIL_FROM"]
    message["To"] = os.environ["EMAIL_TO"]
    message.set_content(
        f"The Chicago permit lead job found {len(new_df)} new permits.\n"
        "The scored Excel report is attached."
    )

    excel_path = files["excel"]
    message.add_attachment(
        excel_path.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=excel_path.name,
    )

    with smtplib.SMTP(os.environ["SMTP_HOST"], int(os.environ["SMTP_PORT"])) as smtp:
        smtp.starttls()
        smtp.login(os.environ["SMTP_USERNAME"], os.environ["SMTP_PASSWORD"])
        smtp.send_message(message)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate scored Chicago permit leads for ChaproNet.")
    parser.add_argument("--config", default="config.yaml", help="Path to configuration YAML.")
    parser.add_argument("--days", type=int, help="Override lookback_days.")
    parser.add_argument("--start-date", help="Inclusive YYYY-MM-DD.")
    parser.add_argument("--end-date", help="Exclusive YYYY-MM-DD; defaults to tomorrow.")
    parser.add_argument("--no-state", action="store_true", help="Do not use or update deduplication state.")
    return parser.parse_args()


def main() -> int:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
    )
    load_dotenv()
    args = parse_args()

    config_path = Path(args.config).resolve()
    project_root = config_path.parent
    config = load_config(config_path)
    paths = setup_paths(project_root, config)

    end = date.fromisoformat(args.end_date) if args.end_date else date.today() + timedelta(days=1)
    days = args.days or int(config["source"].get("lookback_days", 7))
    start = date.fromisoformat(args.start_date) if args.start_date else end - timedelta(days=days)

    LOGGER.info("Pulling permits from %s through %s (end exclusive).", start, end)
    records = fetch_permits(config, start.isoformat(), end.isoformat())
    df = prepare_dataframe(records, config)

    if df.empty:
        LOGGER.warning("No permits matched the selected dates and filters.")
        # Still create valid empty reports with core columns.
        df = pd.DataFrame(columns=[
            "lead_score", "priority", "issue_date", "permit_number", "permit_type",
            "address", "zip_code", "reported_cost", "work_description",
            "score_reasons", "recommended_action", "source_url",
        ])

    permit_ids = set(df["permit_number"].dropna().astype(str)) if "permit_number" in df else set()
    if args.no_state:
        new_df = df.copy()
    else:
        seen = load_seen(paths.seen_file)
        new_df = df[~df["permit_number"].astype(str).isin(seen)].copy()
        save_seen(paths.seen_file, seen | permit_ids)

    files = export_reports(df, new_df, config, paths)
    email_report(files, new_df)

    LOGGER.info("Done. %s total leads; %s new leads.", len(df), len(new_df))
    for label, path in files.items():
        LOGGER.info("%s: %s", label, path)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("Cancelled.", file=sys.stderr)
        raise SystemExit(130)
    except Exception as exc:
        LOGGER.exception("Lead generation failed: %s", exc)
        raise SystemExit(1)
